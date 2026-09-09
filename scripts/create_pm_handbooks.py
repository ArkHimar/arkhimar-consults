from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUTPUT = Path(__file__).resolve().parents[1] / "docs" / "user-guides"
OUTPUT.mkdir(parents=True, exist_ok=True)

BLACK = "000000"
NAVY = "17324D"
PALE = "EDF3F7"
LIGHT = "F5F5F3"
LINE = "D9D9D9"
MUTED = "555555"
ACCENT = "8B4A2B"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def borders(table):
    tblPr = table._tbl.tblPr
    node = tblPr.find(qn("w:tblBorders"))
    if node is None:
        node = OxmlElement("w:tblBorders")
        tblPr.append(node)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), LINE)
        node.append(el)

def cell_margin(cell, top=110, start=130, bottom=110, end=130):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    hdr = OxmlElement("w:tblHeader")
    hdr.set(qn("w:val"), "true")
    trPr.append(hdr)

def keep(paragraph, next=False):
    pPr = paragraph._p.get_or_add_pPr()
    el = OxmlElement("w:keepNext" if next else "w:keepLines")
    pPr.append(el)

def base_doc(title, subtitle, audience):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(.72)
    sec.bottom_margin = Inches(.72)
    sec.left_margin = Inches(.82)
    sec.right_margin = Inches(.82)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, before, after in (("Title", 30, 0, 14), ("Subtitle", 13, 0, 22), ("Heading 1", 20, 18, 9), ("Heading 2", 14, 13, 6), ("Heading 3", 11, 10, 4)):
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = name != "Subtitle"
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    title_p = doc.add_paragraph(style="Title")
    title_p.add_run(title)
    subtitle_p = doc.add_paragraph(style="Subtitle")
    subtitle_p.add_run(subtitle)
    meta = doc.add_table(rows=3, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    meta.autofit = False
    meta.columns[0].width = Inches(1.65)
    meta.columns[1].width = Inches(5.7)
    for i, (label, value) in enumerate((("Audience", audience), ("Release basis", "ArkHimar PM private beta and implementation phases 0 through 8"), ("Document date", "9 September 2026"))):
        meta.cell(i, 0).text = label
        meta.cell(i, 1).text = value
        shade(meta.cell(i, 0), PALE)
        for run in meta.cell(i, 0).paragraphs[0].runs:
            run.bold = True
        for cell in meta.rows[i].cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_margin(cell)
    borders(meta)
    doc.add_paragraph()
    return doc

def footer(doc, short_title):
    for section in doc.sections:
        p = section.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"ArkHimar PM  |  {short_title}  |  Private beta")
        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED)

def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    keep(p, True)
    return p

def paragraph(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    keep(p)
    return p

def bullets(doc, items, numbered=False):
    style = "List Number" if numbered else "List Bullet"
    for item in items:
        p = doc.add_paragraph(style=style)
        if isinstance(item, tuple):
            p.add_run(item[0]).bold = True
            p.add_run(item[1])
        else:
            p.add_run(item)
        p.paragraph_format.space_after = Pt(3)
        keep(p)

def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, header in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = header
        shade(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
        if widths:
            cell.width = Inches(widths[i])
        cell_margin(cell)
    set_repeat_header(t.rows[0])
    for r_idx, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cells[i].width = Inches(widths[i])
            if r_idx % 2:
                shade(cells[i], LIGHT)
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                for run in p.runs:
                    run.font.size = Pt(8.7)
            cell_margin(cells[i])
    borders(t)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t

def page_break(doc):
    doc.add_page_break()

def save(doc, filename, subject):
    footer(doc, subject)
    props = doc.core_properties
    props.title = subject
    props.subject = subject
    props.author = "ArkHimar Consult"
    props.keywords = "ArkHimar PM, project management, governance, user guide"
    props.comments = "Prepared for ArkHimar PM users"
    path = OUTPUT / filename
    doc.save(path)
    return path

def phase_roadmap():
    doc = base_doc("ArkHimar PM Implementation Phases", "Complete roadmap from foundation to release hardening", "Product owners, delivery teams, administrators, project managers and assurance reviewers")
    paragraph(doc, "This roadmap is the authoritative summary of the nine numbered implementation phases used to build ArkHimar PM. Phase 0 established the product and technical baseline. Phases 1 through 8 progressively added secure project controls, governed decisions, delivery management, reporting, exports and release assurance.")
    paragraph(doc, "The phases describe implementation scope. They do not represent the lifecycle stages of a project managed inside ArkHimar PM. A project may use predictive stages, agile iterations or a hybrid combination regardless of the software implementation phase that introduced each capability.")
    heading(doc, "Phase summary", 1)
    table(doc, ["Phase", "Name", "Primary outcome", "Status"], [
        ("0", "Discovery and production baseline", "Repository audit, routes, visual language, deployment model and product risk boundary", "Complete"),
        ("1", "Secure foundation and project setup", "Authentication, tenant workspaces, roles, project creation, overview and audit history", "Complete"),
        ("2", "Initiation and authorization", "Versioned business cases and charters with governed submission and approval", "Complete"),
        ("3", "Core planning", "Master plan, subsidiary plans, requirements, traceability, WBS and scope baseline", "Complete"),
        ("4", "Schedule and cost control", "Network schedule, critical path, cost records, EVM and approved baselines", "Complete"),
        ("5", "Project controls", "Risks, issues, stakeholders, engagement actions and integrated change control", "Complete"),
        ("6", "Execution reporting and governance", "Iterations, work items, WIP, delivery evidence, gates and immutable status reports", "Complete in source"),
        ("7", "Controlled project exports", "Authenticated JSON, CSV, TXT and paginated PDF project packs", "Complete in source"),
        ("8", "Release hardening", "Accessibility, responsive behavior, regression coverage and production validation", "Complete"),
    ], [0.55, 1.65, 4.25, 1.0])
    page_break(doc)
    phases = [
        ("Phase 0 Discovery and production baseline", "Establish what already exists and define a safe path to production.", ["Audit the repository, public routes, deployment settings and current product claims.", "Preserve the ArkHimar architectural visual language across the Consult and PM experiences.", "Identify security, privacy, data, compliance and operational boundaries before adding features.", "Create a phased delivery sequence so each later control builds on stable foundations."], "A reviewed baseline, a release sequence and explicit claims that distinguish implemented capability from planned capability."),
        ("Phase 1 Secure foundation and project setup", "Create the secure place in which all project records live.", ["Email and password authentication, verification and password recovery.", "Tenant-scoped workspaces, membership roles and row-level database isolation.", "Six-step project creation covering identity, strategy, dates, governance, team and template.", "Project overview, transparent readiness checklist, private documents and append-only audit events.", "Workspace team administration, invitations, MFA and session controls."], "A user can establish a workspace, create a project and know that project data is separated by workspace and role."),
        ("Phase 2 Initiation and authorization", "Help an organization decide whether a project should proceed and formally authorize it.", ["Structured business-case narrative, options analysis and deterministic financial calculations.", "Weighted option scoring using fixed and visible criteria.", "Structured project charter covering purpose, objectives, scope, authority, milestones and approval conditions.", "Draft, submission, change request, approval and revision workflows.", "Immutable approved versions and durable decision history."], "The sponsor can trace why the investment was selected and what authority the project manager received."),
        ("Phase 3 Core planning", "Translate authorization into a controlled definition of the work.", ["Master project management plan and 25 subsidiary plan types.", "Requirements register and traceability from objective through deliverable, WBS, activity and evidence.", "Hierarchical WBS and detailed WBS dictionary records.", "Scope statement, acceptance criteria, assumptions, constraints, boundaries and interfaces.", "Approved scope baseline with controlled revisions linked to authorized changes."], "The team has a testable scope and can explain how every requirement becomes planned and accepted work."),
        ("Phase 4 Schedule and cost control", "Establish when work should happen, how much it should cost and how performance will be measured.", ["Activities, milestones, calendars and finish-to-start, start-to-start, finish-to-finish and start-to-finish dependencies.", "Critical-path calculation, cycle rejection, float and bounded Gantt presentation.", "Cost accounts, estimates, commitments, actuals, accruals, forecasts and cash-flow periods.", "Earned value measures including CV, SV, CPI, SPI, EAC, ETC, VAC and TCPI.", "Approved schedule and cost baselines with change-linked revision workflows."], "The organization has approved reference points against which time and cost performance can be assessed."),
        ("Phase 5 Project controls", "Create the registers and decisions required to control uncertainty and change.", ["Risk management plan and threat or opportunity register with probability and impact scoring.", "Issue log with ownership, severity, escalation, root cause and resolution.", "Stakeholder register, power-interest view, engagement targets and actions.", "Private project-management notes protected from ordinary members and historical snapshots.", "Change requests moving from draft through impact analysis, CCB decision, implementation, verification and closure."], "Risks, issues, stakeholders and approved changes remain connected to the project and its baselines."),
        ("Phase 6 Execution reporting and governance", "Manage delivery across predictive, agile and hybrid approaches while preserving formal oversight.", ["Delivery settings for cadence, WIP limit, definition of ready, definition of done and control ceremonies.", "Sprints, phases, releases, stages and continuous-flow periods with goals, capacity and review outcomes.", "Work items linked to requirements and WBS, with ownership, estimates, progress, acceptance criteria, evidence and blockers.", "Governance reviews for stage gates, sprint reviews, releases, steering committees and benefits reviews.", "Immutable status reports that capture delivery, risk, issue, change and earned-value measures at the reporting cut-off."], "Delivery teams can iterate while sponsors retain explicit decision points and reliable evidence."),
        ("Phase 7 Controlled project exports", "Make governed project information portable without weakening access controls.", ["Authenticated, tenant-authorized server-side export generation.", "JSON backup for structured portability and vendor-lock-in avoidance.", "CSV register export for analysis and controlled handoff.", "Readable TXT and paginated PDF project packs.", "Deterministic filenames, private no-store responses and audit events for sensitive exports."], "Authorized users can take a defensible project record outside the platform in standard formats."),
        ("Phase 8 Release hardening", "Verify that the integrated product is usable, resilient and accurately described.", ["Keyboard focus, skip navigation, reduced-motion behavior and responsive layouts.", "Mobile-safe navigation and horizontally bounded delivery views.", "Loading, empty, error and permission states across major workflows.", "Automated regression coverage for routes, security contracts, calculations, execution metrics and exports.", "Production build verification and Lighthouse accessibility and best-practices checks."], "The release passes its documented checks and remains transparent about operational and enterprise limitations."),
    ]
    for title, purpose, outputs, outcome in phases:
        heading(doc, title, 1)
        paragraph(doc, purpose)
        heading(doc, "Delivered capability", 2)
        bullets(doc, outputs)
        paragraph(doc, f"Completion outcome. {outcome}")
    heading(doc, "Current operational boundary", 1)
    paragraph(doc, "The completed phases form a coherent private-beta product. They do not claim certification or full enterprise maturity. The final Phase 6 database migration must be applied before execution, governance and governed reporting mutations are enabled in production. Specialist quality, resource and procurement registers, high-volume resource leveling, threaded comments, meeting conversion, portfolio dashboards, notifications, secure AI, XLSX/PPTX exports, SSO and SCIM remain future extensions unless separately implemented and verified.")
    heading(doc, "How to use this roadmap", 1)
    bullets(doc, ["Use the roadmap to understand product delivery history and release scope.", "Use the role handbooks to operate projects inside the current product.", "Use the product requirements as the forward-looking capability register, because some requirements intentionally remain beyond the private-beta boundary."])
    return save(doc, "ArkHimar_PM_Implementation_Phases.docx", "ArkHimar PM Implementation Phases")

def pm_handbook():
    doc = base_doc("ArkHimar PM Project Manager Handbook", "A practical guide to managing projects and operating the platform", "Project managers, delivery leads, consultants and people new to formal project management")
    paragraph(doc, "This handbook explains how to take a project from an initial business need through authorization, planning, delivery, control, reporting and handover using ArkHimar PM. It also explains the management habits behind each screen so that a new project manager can use the platform with professional judgment rather than treating it as a form-filling exercise.")
    paragraph(doc, "The project manager owns integration. Functional specialists may own estimates, designs, risks or deliverables, but the project manager must keep the business case, charter, scope, schedule, cost, controls, execution data and decisions consistent.")
    heading(doc, "How to use this handbook", 1)
    bullets(doc, ["Read Project management foundations before creating your first live project.", "Follow the lifecycle chapters in sequence for a new project.", "Use the weekly operating rhythm and checklists during delivery.", "Consult the platform boundaries before promising a workflow to sponsors or external parties."])
    heading(doc, "Project management foundations", 1)
    paragraph(doc, "A project is a temporary effort undertaken to create a defined result. It differs from operations, which repeat ongoing work. Effective project management makes five things explicit: why the project exists, what result is authorized, how the work will be delivered, how performance and uncertainty will be controlled, and who may make each decision.")
    table(doc, ["Concept", "Plain meaning", "Why it matters"], [
        ("Business case", "The evidence and options supporting an investment decision", "Prevents a team from delivering work that no longer has sufficient value"),
        ("Charter", "The formal authorization for the project and the project manager", "Defines objectives, boundaries, authority and sponsor accountability"),
        ("Requirement", "A condition or capability the result must satisfy", "Turns stakeholder needs into something that can be traced and tested"),
        ("WBS", "A hierarchical breakdown of the complete project scope", "Creates manageable work packages and prevents omitted work"),
        ("Baseline", "The approved scope, schedule or cost reference point", "Makes variance meaningful and prevents silent changes"),
        ("Risk", "An uncertain event that may affect objectives", "Requires anticipation and a planned response"),
        ("Issue", "A problem that has already happened", "Requires resolution, ownership and often escalation"),
        ("Change request", "A proposed alteration to an approved commitment", "Allows impact analysis and authorization before implementation"),
        ("Governance", "The decision rights, review forums and escalation rules around the project", "Ensures material decisions are made by the right people"),
    ], [1.25, 2.7, 3.55])
    heading(doc, "Your accountability as project manager", 2)
    bullets(doc, [
        ("Integrate the plan. ", "Connect objectives, requirements, WBS, activities, costs, risks, work items and acceptance evidence."),
        ("Maintain reliable records. ", "Update facts at an agreed cut-off and preserve approved versions instead of rewriting history."),
        ("Make decisions visible. ", "Record the owner, options, impact, decision, comments and date."),
        ("Escalate early. ", "Raise matters before tolerance is exceeded, not after a missed commitment becomes unavoidable."),
        ("Protect trust. ", "Separate facts from forecasts, disclose uncertainty and limit access to confidential information."),
    ])
    heading(doc, "Choosing a delivery approach", 2)
    table(doc, ["Approach", "Use when", "Control emphasis"], [
        ("Predictive", "Scope is sufficiently understood and work follows dependent stages", "WBS, network schedule, baselines, milestones and stage gates"),
        ("Agile", "The solution will evolve through frequent feedback and usable increments", "Prioritized backlog, short iterations, WIP, review, retrospective and definition of done"),
        ("Hybrid", "Some commitments need formal baselines while parts of the solution require iteration", "Baseline the external commitments and govern iterative delivery within those boundaries"),
    ], [1.0, 3.2, 3.3])
    page_break(doc)
    heading(doc, "Access and workspace setup", 1)
    bullets(doc, [
        "Sign in with your verified email address. Complete password recovery or the authenticator challenge when required.",
        "Create a workspace only when you are authorized to establish its data boundary. A workspace separates projects, members, documents and audit events.",
        "Open Team and Access to confirm roles. Owners and admins control high-authority decisions. Project managers maintain project records. Members have limited access in the current release.",
        "Use the account security panel to enroll an authenticator and review sessions. Never share credentials or one-time access links.",
    ])
    heading(doc, "Create the project", 1)
    paragraph(doc, "Select New project and complete the six setup steps. Use stable names and codes because they appear in records and exported filenames. Write the problem as the present condition, and write the expected outcome as the measurable future condition.")
    table(doc, ["Setup step", "What to enter", "Quality check"], [
        ("Identity", "Title, short code, description, industry, delivery approach, region, timezone, currency and confidentiality", "A reader can distinguish the project from a department or recurring service"),
        ("Strategic context", "Problem, outcome, objectives, success criteria, benefits and mandatory drivers", "Criteria are measurable and linked to the business need"),
        ("Delivery setup", "Start, finish, complexity, priority, governance model and reporting frequency", "Dates and governance reflect risk and decision speed"),
        ("Team", "Sponsor, project manager, core team and approval roles", "Every key decision has an accountable role"),
        ("Template", "Construction, predictive, agile, hybrid, development or blank", "The template matches the work and will not overwrite future project data"),
        ("Review", "Confirm all setup information", "Dates are logical and no placeholder text remains"),
    ], [1.15, 3.75, 2.6])
    paragraph(doc, "After creation, use the Overview readiness checklist as a guide, not as proof that the project is healthy. A high readiness score means that expected records exist. It does not mean the assumptions are correct or delivery is on track.")
    heading(doc, "Develop and govern the business case", 1)
    bullets(doc, [
        "Describe the context, problem or opportunity, strategic alignment, intended outcomes, funding assumptions, risks, dependencies, constraints and sustainability implications.",
        "Add at least two genuine options. Include a do-minimum or defer option where appropriate. Score each option using the visible weighted criteria and explain the recommendation.",
        "Enter estimated costs, benefits, discount rate and period cash flows carefully. ROI, NPV and payback are calculations based on your assumptions; they do not validate those assumptions.",
        "Save the draft, review it with finance and subject specialists, then submit it. The submitted version locks while an owner or admin decides.",
        "If changes are requested, address the comments and resubmit. Once approved, create a revision rather than editing the approved version in place.",
    ])
    heading(doc, "Create and approve the charter", 1)
    paragraph(doc, "The charter converts an approved idea into an authorized project. Keep it concise enough for the sponsor to understand while including the boundaries needed for control.")
    bullets(doc, ["State purpose, measurable objectives and success criteria.", "Define high-level scope, requirements, deliverables, exclusions, assumptions and constraints.", "Record key milestones, budget, risks and stakeholders.", "Define the project manager's authority, approval requirements and exit criteria.", "Submit the charter only when sponsor, scope, authority and success criteria are clear."])
    heading(doc, "Plan scope requirements and the WBS", 1)
    paragraph(doc, "Planning should make the complete authorized work visible. Begin with outcomes and acceptance, then decompose the work. Do not begin by typing a long task list with no link to requirements or deliverables.")
    bullets(doc, [
        "Write the project scope, product scope, deliverables, inclusions, exclusions, acceptance criteria, assumptions, constraints, boundaries and interfaces.",
        "Complete the master plan and the subsidiary plans that are proportionate to the project. Record methods, roles, thresholds, reporting, escalation and approvals.",
        "Create requirements with a source, owner, priority, acceptance criteria and verification method. Trace each requirement to an objective, benefit, deliverable, WBS item, activity and evidence where applicable.",
        "Build the WBS around deliverables. Decompose until a work package can be estimated, assigned and accepted. Use the dictionary to explain scope, exclusions, assumptions, resources, quality criteria, dependencies and acceptance.",
        "Resolve orphan requirements and ambiguous work packages before requesting the scope baseline.",
        "An owner or admin approves the baseline. Later changes must cite an authorized change request before a revision begins.",
    ])
    heading(doc, "Build and control the schedule", 1)
    paragraph(doc, "Convert WBS work packages into activities. Define dependencies based on how work must flow, not on the order in which rows were entered.")
    bullets(doc, [
        "Assign a unique activity ID, WBS link, name, duration, owner and calendar. Mark true milestones with zero duration.",
        "Use FS when a successor starts after its predecessor finishes; SS when starts are linked; FF when finishes are linked; and SF only when the unusual relationship is genuinely required.",
        "Use lead or lag sparingly and explain it. Hidden waiting time makes a schedule difficult to defend.",
        "Review critical activities and float. A critical activity has no scheduling flexibility under the current network assumptions.",
        "During delivery, record actual starts, actual finishes, remaining duration and percent complete. Forecast from remaining work rather than preserving obsolete dates.",
        "Baseline only after scope, logic, resources and dates have been reviewed. Use a change-linked revision when an approved commitment changes.",
    ])
    heading(doc, "Build and control the budget", 1)
    bullets(doc, [
        "Set the budget at completion and distinguish contingency reserve from management reserve.",
        "Create cost records with cost account, WBS/activity link, type, classification, vendor, amount, tax, due date and invoice status.",
        "Maintain cash-flow periods and distinguish planned value, earned value and actual cost.",
        "Interpret CPI below 1 as cost inefficiency and SPI below 1 as less value completed than planned. Check the underlying data before escalating a ratio.",
        "Review EAC, ETC, VAC and TCPI. Explain which forecast assumption applies and do not use a single formula mechanically when future performance will differ from the past.",
        "Approve the cost baseline only after scope, estimates, reserves, funding and timing are consistent.",
    ])
    heading(doc, "Manage risks issues stakeholders and changes", 1)
    table(doc, ["Control", "Minimum management action", "Escalate when"], [
        ("Risk", "State cause, uncertain event, impact, probability, impact score, owner, response, trigger and due date", "Exposure, proximity or response delay exceeds the agreed threshold"),
        ("Issue", "Record the current problem, owner, severity, priority, due date, root cause and resolution", "The owner cannot resolve it within authority or tolerance"),
        ("Stakeholder", "Assess influence, interest, current engagement, desired engagement, expectations and concerns", "Resistance or missing support threatens an objective or decision"),
        ("Change", "Describe the reason and assess scope, schedule, cost, quality, resources, risks, contracts, benefits and security", "An approved baseline, commitment or benefit may change"),
    ], [1.0, 4.25, 2.25])
    paragraph(doc, "Never use a risk record for a problem that has already occurred. Never implement a material baseline change while its request is still under review. The current change workflow is Draft, Submitted, Impact analysis, CCB review, decision, Implementation, Verification and Closed. Owner or admin authority is required for approval decisions and closure after verification.")
    heading(doc, "Run predictive agile and hybrid delivery", 1)
    bullets(doc, [
        "Open Delivery and Governance. Set the delivery method, cadence, WIP limit, reporting cut-off, definition of ready, definition of done and control ceremonies.",
        "Create the relevant delivery periods: sprints, phases, releases, stages or continuous-flow periods. Give each a clear goal, dates and realistic capacity.",
        "Create work items and link them to requirements and WBS records. Assign an owner, priority, estimate, status and due date.",
        "Do not start work that fails the definition of ready. Limit active work so the team finishes valuable items before starting more.",
        "Record acceptance criteria before delivery and evidence before marking work Done. Percent complete should match observable work, not optimism.",
        "Capture review outcomes and retrospective lessons. Convert lessons into changes to the next plan or iteration.",
    ])
    heading(doc, "Use governance reviews", 2)
    paragraph(doc, "Create a governance review for a stage gate, sprint review, release review, steering committee or benefits review. Define the decision criteria before the meeting. Add evidence and a recommendation, then move the review from Draft to Scheduled and Ready. An owner or admin records approval, rejection or deferral with required comments. Approved reviews may then be closed.")
    heading(doc, "Publish reliable status reports", 1)
    paragraph(doc, "Choose a reporting cut-off and update the source records before writing the narrative. ArkHimar PM captures delivery completion, blocked work, risks, issues, changes, PV, EV, AC, CPI and SPI when the report is published. Published reports are immutable.")
    bullets(doc, [
        "State overall health as On track, Watch or At risk. Explain any judgment that differs from the calculated indicators.",
        "Write the executive summary around outcomes, variance and action. Avoid listing activity without explaining its effect on objectives.",
        "Record achievements, next priorities, forecast milestones, exceptions and decisions needed.",
        "Name the decision owner and required date in the Decisions needed section.",
        "Correct a published error through a new report and an explanatory record. Do not attempt to erase the original snapshot.",
    ])
    heading(doc, "Manage communications documents and access", 1)
    bullets(doc, [
        "Use Communications to prepare branded email or letter content and select only authorized private attachments. Confirm recipients and classification before sending.",
        "Use controlled documents for records that need a document ID, version, review, approval, archive status or controlled export. Approved versions are immutable.",
        "Use private file attachments for working files. A signed access URL lasts briefly and should not be forwarded casually.",
        "Create external shares only for approved controlled-document versions. Set a suitable expiry and download limit, copy the token once, and revoke it when access is no longer required.",
        "Use Team and Access to apply least privilege. Remove access when a person leaves the project or no longer needs the workspace.",
        "Use the Integration Hub only with approved systems. API keys are shown once, stored as hashes and should be revoked if exposed.",
    ])
    heading(doc, "Corporate project operating rhythm", 1)
    table(doc, ["Cadence", "Project manager activity", "Expected output"], [
        ("Daily or frequent", "Review blockers, critical work, urgent issues and decision deadlines", "Updated owners and immediate action"),
        ("Weekly", "Update schedule, cost, RAID, changes, stakeholder actions and execution board", "Reliable control data at the agreed cut-off"),
        ("Fortnightly or sprint", "Plan work, review completed evidence and hold a retrospective", "Accepted increment and improved next plan"),
        ("Monthly", "Reforecast completion, prepare status, review benefits and chair governance", "Published report and recorded decisions"),
        ("At each gate", "Confirm criteria, artifacts, risks, affordability and readiness", "Approve, reject or defer with comments"),
    ], [1.15, 4.2, 2.15])
    heading(doc, "Meetings and decision discipline", 2)
    paragraph(doc, "For each meeting, circulate the purpose and agenda, identify the decision owner, record attendees, capture decisions and assign actions with due dates. ArkHimar PM does not yet provide a normalized meeting record or note-to-action conversion. Maintain the meeting record as a controlled document or approved corporate record, then update the relevant risk, issue, change, work item or governance review in ArkHimar PM.")
    heading(doc, "Escalation practice", 2)
    paragraph(doc, "Escalation is a request for a decision or support that exceeds the project manager's authority. State the condition, evidence, effect on objectives, options, recommendation, decision owner and latest useful decision date. Do not escalate a vague concern, and do not conceal an exceeded tolerance because the team hopes to recover later.")
    heading(doc, "First thirty days checklist", 1)
    bullets(doc, [
        "Confirm sponsor, project manager, decision rights, confidentiality and reporting cadence.",
        "Validate the business need, outcomes, options, affordability and major assumptions.",
        "Obtain an approved charter before committing substantial delivery resources.",
        "Identify stakeholders and agree how decisions, risks and issues will be communicated.",
        "Create requirements, acceptance criteria, WBS, schedule and cost plan.",
        "Resolve planning inconsistencies and obtain appropriate baselines.",
        "Set the execution method, WIP, definitions and governance reviews.",
        "Establish the weekly update cut-off and publish the first status report.",
        "Confirm document controls, access rights, external-sharing rules and audit expectations.",
    ])
    heading(doc, "Common failure modes", 1)
    table(doc, ["Failure", "Why it is harmful", "Corrective action"], [
        ("Treating the readiness score as project health", "Records may exist but contain weak assumptions", "Review evidence, forecast and exceptions separately"),
        ("Marking work complete without evidence", "Acceptance becomes subjective and disputes increase", "Define criteria first and attach or cite completion evidence"),
        ("Updating a baseline to match actual performance", "Variance and decision history disappear", "Preserve the baseline and use a governed change revision"),
        ("Reporting activity instead of outcomes", "Sponsors cannot see value, variance or decisions", "Explain achievement, forecast, risk and required action"),
        ("Assigning every risk to the PM", "Responses lack accountable subject expertise", "Assign the person with authority and capability to act"),
        ("Starting too much work", "Queues grow and completion slows", "Apply the WIP limit and finish priority items"),
    ], [1.65, 2.85, 3.0])
    heading(doc, "Current platform boundaries", 1)
    paragraph(doc, "The private beta currently supports the governed records described above after the final database migration is applied. Threaded comments, mentions, watchers, notification preferences, normalized meeting records, team-member self-service work updates, portfolio dashboards, specialist quality/resource/procurement registers, PPTX/XLSX exports, SSO/SCIM and secure AI assistance are not yet production capabilities. Use approved corporate processes alongside ArkHimar PM where these functions are required.")
    heading(doc, "Glossary", 1)
    table(doc, ["Term", "Meaning"], [
        ("BAC", "Budget at completion, the approved total performance-measurement budget"),
        ("CCB", "Change Control Board, the authorized group that decides material changes"),
        ("CPI", "Cost Performance Index, earned value divided by actual cost"),
        ("EAC", "Estimate at completion, the forecast total cost"),
        ("EVM", "Earned Value Management, integrated measurement of scope, schedule and cost"),
        ("RACI", "Responsible, Accountable, Consulted and Informed role assignment"),
        ("RAG", "Red, amber and green or equivalent health classification"),
        ("SPI", "Schedule Performance Index, earned value divided by planned value"),
        ("Tolerance", "The approved variance range within which management may act without escalation"),
        ("WIP", "Work in progress, the number of started items not yet completed"),
        ("WBS", "Work Breakdown Structure, the hierarchical decomposition of total project scope"),
    ], [1.5, 6.0])
    return save(doc, "ArkHimar_PM_Project_Manager_Handbook.docx", "ArkHimar PM Project Manager Handbook")

def sponsor_handbook():
    doc = base_doc("ArkHimar PM Sponsor Handbook", "How to authorize direct and govern projects", "Executive sponsors, investment owners, steering committee members and senior accountable owners")
    paragraph(doc, "This handbook explains how a sponsor uses ArkHimar PM to authorize a project, set decision boundaries, review performance and protect the intended business outcome. It assumes no prior project-management training.")
    paragraph(doc, "The sponsor owns the business need, continued justification and senior decisions. The project manager organizes and controls delivery. A sponsor should not run the daily plan, but must remain available when authority, funding, priorities or organizational support are required.")
    heading(doc, "Sponsor responsibilities", 1)
    table(doc, ["Responsibility", "What good sponsorship looks like"], [
        ("Own the outcome", "Confirm that the problem, outcome, benefits and success measures remain relevant"),
        ("Authorize the project", "Approve a credible business case and charter with clear authority and boundaries"),
        ("Provide decisions", "Decide material changes and governance reviews before delay removes useful options"),
        ("Secure support", "Resolve cross-functional priorities, funding and senior stakeholder barriers"),
        ("Protect integrity", "Require accurate reporting, evidence and immutable decision history"),
        ("Confirm closure and benefits", "Accept the delivered outcome and ensure benefits ownership continues after project closure"),
    ], [2.0, 5.5])
    heading(doc, "Access the platform safely", 1)
    bullets(doc, ["Accept only an invitation sent to your authorized email address.", "Use your own account. Do not approve through another person's session.", "Enroll an authenticator when required and review account sessions periodically.", "Treat project classification seriously. Do not forward controlled exports or external-share links without authorization.", "Your ability to approve depends on your workspace role. Owner and admin roles hold the current high-authority decision permissions."])
    heading(doc, "Review the business case", 1)
    paragraph(doc, "Open Business Case and review the submitted version. The decision is whether the proposed investment remains the best available response to the business need, not whether the document is attractively written.")
    bullets(doc, [
        "Confirm the problem or opportunity is evidenced and material.",
        "Check that strategic alignment and intended benefits name observable outcomes.",
        "Compare the options, including affordability, delivery feasibility, risk and a credible alternative to immediate investment.",
        "Challenge cost, benefit, discount-rate and cash-flow assumptions. ROI, NPV and payback are only as reliable as their inputs.",
        "Check dependencies, constraints, sustainability impacts and major risks.",
        "Approve when the recommended option is justified. Request changes when evidence or assumptions are inadequate. Record comments clearly enough for the team to act.",
    ])
    heading(doc, "Review the charter", 1)
    bullets(doc, [
        "Confirm the purpose and objectives match the approved business case.",
        "Check that success criteria are measurable and owned.",
        "Review scope, exclusions, deliverables, major milestones and budget.",
        "Confirm the named project manager has sufficient authority and knows the escalation thresholds.",
        "Confirm approval roles, exit criteria and sponsor responsibilities.",
        "Approve the submitted version only when the organization is ready to commit. An approved version is immutable; later changes require a revision.",
    ])
    heading(doc, "Understand baselines and change control", 1)
    paragraph(doc, "A baseline is the approved reference for scope, schedule or cost. It creates accountability because current forecasts can be compared with the commitment. Do not ask the project manager to rewrite a baseline merely to make performance appear favorable.")
    table(doc, ["Decision", "Sponsor questions"], [
        ("Scope baseline", "Is the required outcome complete, testable and free from material ambiguity?"),
        ("Schedule baseline", "Is the logic credible, are critical milestones achievable, and are key dependencies owned?"),
        ("Cost baseline", "Does funding cover authorized scope, realistic estimates and appropriate reserves?"),
        ("Change request", "What changes, what is the quantified effect, what options exist, and who accepts the consequences?"),
    ], [1.6, 5.9])
    paragraph(doc, "In ArkHimar PM, a change proceeds through drafting, submission, impact analysis and CCB review before a decision. Owners or admins approve, reject or defer. Approved changes move through implementation, verification and closure. Ask for evidence that the decision was implemented and that related baselines were revised where necessary.")
    heading(doc, "Use governance reviews", 1)
    paragraph(doc, "Governance reviews provide formal points to continue, change, pause or stop work. Open Delivery and Governance and examine the review type, criteria, evidence and recommendation. A review marked Ready requires a decision from an owner or admin.")
    bullets(doc, ["Approve when criteria are met and the residual exposure is acceptable.", "Reject when continuation is not justified or the proposed outcome is no longer required.", "Defer when a specific piece of evidence or action is required before decision. State exactly what is missing and the next review date.", "Close an approved review after its required follow-up is complete."])
    heading(doc, "Read a status report", 1)
    paragraph(doc, "Start with the outcome and forecast, then investigate the indicators. A report is useful when it tells you whether objectives remain achievable, what changed since the last cut-off and which decisions require your authority.")
    table(doc, ["Indicator", "Interpretation", "Question to ask"], [
        ("Overall health", "Management judgment: On track, Watch or At risk", "What evidence supports this rating and what would change it?"),
        ("Delivery completion", "Completed estimated work compared with total estimated work", "Is completed work accepted and supported by evidence?"),
        ("Blocked work", "Items that cannot progress", "Who can remove the blocker and by when?"),
        ("CPI", "Below 1 suggests the project is earning less value per unit of cost than planned", "Is this temporary, and what is the revised cost forecast?"),
        ("SPI", "Below 1 suggests less value has been completed than planned", "Which milestone or benefit is affected?"),
        ("Critical risks", "High probability-impact exposure requiring attention", "Is the response funded, owned and due before the trigger?"),
        ("Decisions needed", "Matters outside the PM's authority", "What happens if the decision is late?"),
    ], [1.3, 3.15, 3.05])
    paragraph(doc, "Published reports are immutable snapshots. If an error is later discovered, require a corrected report and explanation rather than deleting the historical record.")
    heading(doc, "Sponsor meeting rhythm", 1)
    bullets(doc, [
        "Hold a short regular sponsor-PM meeting focused on forecast, exceptions and decisions, not task-by-task narration.",
        "Attend stage gates, steering committees and major release reviews where your authority is required.",
        "Review the business case when cost, timing, benefits, regulation or strategy changes materially.",
        "Respond to escalations by the latest useful decision date. Silence is itself a project constraint.",
        "Confirm stakeholder support and remove organizational barriers that the project manager cannot resolve.",
    ])
    heading(doc, "Questions that improve decisions", 1)
    bullets(doc, ["What has changed since the last approved position?", "Which assumption is least certain?", "What is the effect on outcome, benefit, date and total cost?", "What options were considered, including stopping or reducing scope?", "What decision is required from me, and by what date?", "What evidence will show that the decision was implemented?", "Who owns the benefit after project handover?"])
    heading(doc, "Warning signs", 1)
    table(doc, ["Warning sign", "Sponsor response"], [
        ("All reports remain On track despite repeated missed milestones", "Ask for source data, current forecast and the health-rating rationale"),
        ("A change is already implemented before approval", "Stop unauthorized work where practical and require impact assessment"),
        ("Risks have no owners or overdue responses", "Assign accountable owners and confirm resources"),
        ("Completion rises but acceptance evidence is missing", "Separate work performed from work accepted"),
        ("The PM escalates the same dependency repeatedly", "Use sponsor authority to resolve the organizational constraint"),
        ("The expected benefit is no longer credible", "Reassess the business case and consider pause, redesign or termination"),
    ], [3.25, 4.25])
    heading(doc, "Closure and benefits", 1)
    paragraph(doc, "Before authorizing closure, confirm deliverable acceptance, outstanding defects or obligations, operational handover, document retention, financial closure, lessons and a named benefits owner. ArkHimar PM does not yet provide a normalized closure and benefits module. Maintain the formal closure record as a controlled document and preserve follow-up governance outside the platform until those capabilities are implemented.")
    heading(doc, "Current platform boundaries", 1)
    paragraph(doc, "The current private beta records approvals and governance within the main project views; it does not yet provide a separate notification-driven sponsor inbox or focused approval page. Portfolio dashboards, PPTX executive packs, threaded comments, meeting records and automated reminders remain planned. Use your organization's approved communication and governance channels alongside the platform.")
    return save(doc, "ArkHimar_PM_Sponsor_Handbook.docx", "ArkHimar PM Sponsor Handbook")

def stakeholder_handbook():
    doc = base_doc("ArkHimar PM Stakeholder Handbook", "How team members reviewers clients and other stakeholders work with project information", "Internal team members, functional managers, reviewers, clients, consultants, suppliers and external stakeholders")
    paragraph(doc, "This handbook explains how people other than the project manager or sponsor participate in an ArkHimar PM project. It covers safe access, clear updates, risk and issue reporting, evidence, reviews and professional conduct in a corporate project setting.")
    paragraph(doc, "A stakeholder is any person or group that can affect the project, is affected by it or believes it may be affected. Stakeholders do not all need the same access. The project manager and workspace administrators should grant only the information and actions required for each role.")
    heading(doc, "Know your participation role", 1)
    table(doc, ["Role", "Typical contribution", "Expected behavior"], [
        ("Team member", "Completes assigned work and provides progress and evidence", "Raise blockers early and distinguish finished work from accepted work"),
        ("Functional manager", "Provides specialist people, standards and assurance", "Resolve resource conflicts and validate technical estimates"),
        ("Reviewer or approver", "Assesses an artifact or decision against criteria", "Review the submitted version and record actionable comments"),
        ("Client or user representative", "Clarifies needs and validates outcomes", "Give timely feedback and confirm acceptance against agreed criteria"),
        ("Supplier or consultant", "Delivers contracted work or advice", "Use agreed interfaces, protect information and disclose variance"),
        ("External observer", "Receives a permitted approved record", "Use the link only for its intended purpose and do not redistribute it"),
    ], [1.45, 3.0, 3.05])
    heading(doc, "Access and confidentiality", 1)
    bullets(doc, [
        "Accept a workspace invitation only through the intended email account. The invitation is time-limited and bound to that address.",
        "Use your own account and complete any authenticator challenge. Report unexpected sessions or access immediately.",
        "Read the project confidentiality label and follow your organization's handling rules.",
        "Do not assume workspace membership permits access to every detail. Private project-manager stakeholder notes and some administrative records are restricted.",
        "External shares should point only to approved controlled-document versions. Links can expire, reach a download limit or be revoked.",
        "Do not upload personal, commercially sensitive or regulated information unless the organization has approved the platform and handling process for that data.",
    ])
    heading(doc, "Understand the project before acting", 1)
    bullets(doc, ["Read the Overview to understand the problem, outcome, sponsor, project manager, dates, approach and current readiness.", "Read the approved charter for objectives, scope, exclusions, authority and success criteria.", "Check the relevant requirement, WBS item, activity or work item before beginning work.", "Confirm the current approved version. A downloaded or emailed copy may have been superseded.", "Ask the project manager when a responsibility, due date, acceptance criterion or decision right is unclear."])
    heading(doc, "Give a professional progress update", 1)
    paragraph(doc, "A useful update is factual, dated and connected to a commitment. It allows the project manager to forecast and act. Report what is complete, what remains, what evidence exists and what prevents progress.")
    table(doc, ["Include", "Example"], [
        ("Reference", "WI-014 or ACT-032"),
        ("Status at cut-off", "In progress at 9 September 2026"),
        ("Completed work", "Draft reviewed by Legal and comments incorporated"),
        ("Remaining work", "Finance validation and sponsor acceptance"),
        ("Forecast", "Ready for review on 12 September if cost data arrives by 10 September"),
        ("Blocker", "Supplier has not provided the required certificate"),
        ("Evidence", "Approved drawing reference, test result, signed acceptance or controlled-document version"),
        ("Support needed", "Procurement lead to escalate the certificate by 10 September"),
    ], [1.75, 5.75])
    paragraph(doc, "Do not report 90 percent complete for several weeks without defining remaining work. Percent complete must reflect observable progress. A work item marked Done should satisfy its acceptance criteria and cite evidence.")
    heading(doc, "Raise risks issues and changes correctly", 1)
    table(doc, ["Situation", "Use", "What to provide"], [
        ("Something uncertain may happen", "Risk", "Cause, event, potential impact, timing, probability and suggested response"),
        ("A problem has already happened", "Issue", "Facts, current impact, owner needed, urgency and proposed resolution"),
        ("An approved commitment may need alteration", "Change request", "Reason, affected scope/date/cost/quality/risk and available options"),
        ("You cannot continue assigned work", "Blocker on the work item", "Specific dependency, person needed and latest useful resolution date"),
    ], [2.2, 1.35, 3.95])
    paragraph(doc, "Raise matters early. Do not wait for the next status meeting when safety, legal compliance, information security, material cost, critical-path timing or reputation may be affected. Use the organization's urgent escalation route in addition to the platform when immediate action is required.")
    heading(doc, "Participate in agile or hybrid delivery", 1)
    bullets(doc, [
        "Understand the iteration or stage goal before accepting work.",
        "Confirm the work item meets the definition of ready: its purpose, owner, inputs and acceptance criteria are clear.",
        "Respect the WIP limit. Finish or unblock current priority work before starting additional items.",
        "Update status honestly: Backlog, Ready, In progress, Review, Blocked, Done or Cancelled.",
        "Bring completed evidence to the review. Feedback may create revised work; it does not rewrite the completed iteration's history.",
        "Use the retrospective to identify a practical process change, then apply it in the next cycle.",
    ])
    heading(doc, "Participate in reviews and approvals", 1)
    paragraph(doc, "Review against stated criteria, not personal preference. Read the current submitted version, examine evidence and identify whether a comment is mandatory, recommended or a question. State the affected section and the action required.")
    bullets(doc, ["Approve only within your delegated authority.", "Do not give informal approval in a meeting if the formal record requires a platform decision.", "When requesting changes, explain what must change and why the current evidence is insufficient.", "Disclose conflicts of interest and do not approve your own work where independent review is required.", "Meet the decision deadline or tell the project manager immediately when more time or evidence is necessary."])
    heading(doc, "Use documents and external shares", 1)
    bullets(doc, [
        "A controlled document has an ID, status and version. Draft is working content; In review awaits assessment; Approved is immutable; Archived is retained but no longer current.",
        "Check the document ID and version before relying on a file.",
        "Use private attachments for authorized working files and approved controlled exports for formal records.",
        "An external share may expose only the approved version and format selected by the sender. Do not modify the URL or attempt to discover other content.",
        "Tell the sender if a share reaches its limit, expires or appears to expose the wrong information. Do not ask someone to bypass the control by forwarding unrestricted copies.",
    ])
    heading(doc, "Corporate communication standards", 1)
    bullets(doc, [
        "Use the project code and record reference in messages and meeting notes.",
        "Separate fact, assumption, forecast, opinion and decision.",
        "State the owner and due date for every action.",
        "Use concise subject lines and put the requested decision near the beginning.",
        "Avoid parallel private records that conflict with the current project record. Send the project manager the approved update for incorporation.",
        "Challenge respectfully with evidence. Escalation should protect the outcome, not assign blame.",
    ])
    heading(doc, "Meeting participation", 1)
    bullets(doc, ["Read the agenda and decision papers before the meeting.", "Arrive prepared to state status, evidence, forecast, blocker and support needed.", "Confirm decisions aloud, including owner, effective date and conditions.", "Review assigned actions after the meeting and correct misunderstandings promptly.", "ArkHimar PM does not yet include a normalized meeting module. Use the approved meeting record, then update the relevant platform registers."])
    heading(doc, "What different access methods allow", 1)
    table(doc, ["Access method", "What it is for", "Important limitation"], [
        ("Workspace owner or admin", "Administration and high-authority decisions", "Grant only to people who require that authority"),
        ("Project manager", "Maintaining project plans, controls, execution and reports", "Responsible for integrated accuracy and escalation"),
        ("Member", "Authorized workspace visibility and permitted file contribution", "Current beta does not yet provide self-service My Work updates"),
        ("Controlled external share", "Time-limited access to one approved document version", "Does not provide project or workspace discovery"),
        ("Exported project pack", "Authorized portable project record", "The recipient must protect the downloaded copy outside ArkHimar PM"),
    ], [1.65, 3.15, 2.85])
    heading(doc, "Current platform boundaries", 1)
    paragraph(doc, "Threaded comments, mentions, reactions, watchers, notification preferences, a My Work page, direct team-member progress updates, normalized meeting records and external commenting or approval are planned product requirements rather than current private-beta capabilities. Until implemented, team members should submit updates through the organization's approved channel and the project manager should enter or govern the resulting project record. External stakeholders currently receive approved documents through controlled shares rather than access to the full project.")
    heading(doc, "Stakeholder checklist", 1)
    bullets(doc, ["I understand the project outcome, my role and my decision authority.", "I know which current version governs my work.", "My progress update includes evidence, remaining work and forecast.", "I have raised risks, issues, blockers and possible changes early.", "I protect confidential information and do not redistribute access links.", "I record approvals only through the authorized process.", "I close actions by the agreed date or escalate before the commitment is missed."])
    return save(doc, "ArkHimar_PM_Stakeholder_Handbook.docx", "ArkHimar PM Stakeholder Handbook")

if __name__ == "__main__":
    paths = [phase_roadmap(), pm_handbook(), sponsor_handbook(), stakeholder_handbook()]
    for path in paths:
        print(path)
